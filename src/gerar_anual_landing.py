"""AB01: gera candidato anual com uma linha por documento válido da 00_landing.

Uso: python -u src/gerar_anual_landing.py . [--run-id ID]
Nunca altera A-OS_2022.xlsx. Um documento com vários itens recebe o total
somado; qtd e valor unit ficam vazios, pois não há um único par representativo.
"""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook

MESES = ("JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO")
COLUNAS = ("mês", "formulário", "prestador solicitado", "tipo", "número de controle",
           "descrição do defeito", "área/equipamento", "local",
           "descrição do serviço realizado", "peças e equipamentos necessários",
           "qtd", "valor unit", "valor total")
CODIGO = re.compile(r"OSM\.MQ/EQ-(\d+)F2", re.I)


def numero(value: object) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError("valor numérico ausente")
    try:
        return Decimal(str(value).replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"valor numérico inválido: {value!r}") from exc


def extrair(arquivo: Path, mes: str) -> tuple[list, dict, list[dict]]:
    wb = load_workbook(arquivo, read_only=True, data_only=True, keep_links=False)
    try:
        if "OSM" not in wb:
            raise ValueError("aba OSM ausente")
        ws = wb["OSM"]
        bruto = str(ws["H2"].value or "").strip()
        m = CODIGO.search(bruto)
        if not m:
            raise ValueError(f"formulário inválido no cabeçalho H2: {bruto!r}")
        formulario = f"OSM.MQ/EQ-{m.group(1)}F2"
        itens = []
        for row in ws.iter_rows(min_row=36, max_row=60, min_col=1, max_col=10):
            qtd, descricao, unit, total = row[0].value, row[1].value, row[7].value, row[9].value
            if total in (None, "", 0) and qtd in (None, "") and unit in (None, ""):
                continue
            if total in (None, "", 0):
                continue
            q, u, t = numero(qtd), numero(unit), numero(total)
            if abs(q * u - t) > Decimal("0.005"):
                raise ValueError(f"fórmula inválida na linha {row[0].row}: {q} × {u} != {t}")
            itens.append({"linha_origem": row[0].row, "peça": descricao, "qtd": str(q),
                          "valor_unit": str(u), "valor_total": str(t)})
        if not itens:
            raise ValueError("nenhum item com valor no quadro A36:J60")
        total = sum((Decimal(i["valor_total"]) for i in itens), Decimal(0))
        candidato = [mes, formulario, ws["A19"].value, ws["A24"].value,
                     ws["A26"].value, ws["C19"].value, ws["F19"].value,
                     ws["H19"].value, ws["A31"].value,
                     itens[0]["peça"] if len(itens) == 1 else " | ".join(
                         str(i["peça"] or f"ITEM LINHA {i['linha_origem']}") for i in itens),
                     float(Decimal(itens[0]["qtd"])) if len(itens) == 1 else None,
                     float(Decimal(itens[0]["valor_unit"])) if len(itens) == 1 else None,
                     float(total)]
        info = {"arquivo": str(arquivo), "formulário": formulario,
                "itens": len(itens), "valor_total": str(total)}
        return candidato, info, itens
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("projeto", nargs="?", type=Path, default=Path("."))
    parser.add_argument("--run-id", default=None)
    args = parser.parse_args()
    raiz = args.projeto.resolve()
    landing = raiz / "datalake" / "00_landing"
    if not landing.is_dir():
        parser.error(f"Pasta ausente: {landing}")
    run_id = args.run_id or datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    qualidade = raiz / "quality" / "AB01" / run_id
    docs = raiz / "docs" / "execucoes" / "AB01" / run_id
    qualidade.mkdir(parents=True, exist_ok=False)
    docs.mkdir(parents=True, exist_ok=False)
    linhas, detalhes, excecoes, ignorados = [], [], [], []
    por_mes = Counter()
    for pasta in sorted(landing.iterdir()):
        if not pasta.is_dir():
            continue
        prefixo = pasta.name[:2]
        if not prefixo.isdigit() or not 1 <= int(prefixo) <= 12:
            continue
        mes = MESES[int(prefixo) - 1]
        for arq in sorted(pasta.rglob("*")):
            if not arq.is_file() or arq.suffix.lower() not in {".xlsm", ".xlsx"}:
                continue
            if arq.name.startswith("~$") or not zipfile.is_zipfile(arq):
                ignorados.append(str(arq.relative_to(raiz)))
                continue
            try:
                linha, info, itens = extrair(arq, mes)
            except Exception as exc:
                excecoes.append({"arquivo": str(arq.relative_to(raiz)), "erro": str(exc)})
                continue
            info["arquivo"] = str(arq.relative_to(raiz))
            info["itens_detalhados"] = itens
            if len(itens) > 1:
                info["regra"] = "qtd e valor unit vazios; valor total é a soma dos itens"
            detalhes.append(info)
            linhas.append(linha)
            por_mes[mes] += 1
        print(f"[AB01] {mes}: {por_mes[mes]} documentos extraídos", flush=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "OS_2022"
    ws.append(COLUNAS)
    for linha in linhas:
        ws.append(linha)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    destino = qualidade / "A-OS_2022_gerado.xlsx"
    wb.save(destino)
    por_codigo = Counter(linha[1] for linha in linhas)
    repetidos = {k: v for k, v in por_codigo.items() if v > 1}
    ambiguos = [v for v in detalhes if v["itens"] > 1]
    status = "APROVADO" if not (excecoes or repetidos or ambiguos) else "REPROVADO"
    relatorio = {"etapa": "AB01", "run_id": run_id, "status": status,
                 "granularidade": "uma linha por arquivo válido da 00_landing",
                 "linhas_geradas": len(linhas), "por_mes": dict(por_mes),
                 "valor_total_gerado": round(sum(linha[-1] for linha in linhas), 2),
                 "arquivos_ignorados": ignorados, "falhas_extracao": excecoes,
                 "codigos_oficiais_repetidos": repetidos, "documentos_multiplos_itens": ambiguos,
                 "linhagem": detalhes, "arquivo_gerado": str(destino.relative_to(raiz)),
                 "decisao": "Candidato para revisão; não substitui o anual existente."}
    (docs / "conclusao_AB01.json").write_text(
        json.dumps(relatorio, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    md = ["# AB01 — anual reconstruído da origem", "", f"Execução: `{run_id}`", "",
          f"Status: **{status}**", "", f"Documentos transformados: **{len(linhas)}**.",
          f"Valor somado: **R$ {relatorio['valor_total_gerado']:.2f}**.",
          f"Falhas de extração: **{len(excecoes)}**; códigos repetidos: **{len(repetidos)}**; "
          f"documentos com vários itens: **{len(ambiguos)}**.", "",
          "Uma linha por documento. Para vários itens, a quantidade e o valor unitário não têm "
          "um único valor: permanecem vazios, e o total é somado. Consulte o JSON para os itens.",
          "O arquivo gerado é candidato e não substitui A-OS_2022.xlsx."]
    (docs / "conclusao_AB01.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"[AB01] {len(linhas)} linhas; R$ {relatorio['valor_total_gerado']:.2f}; "
          f"status {status}; evidências {docs}", flush=True)
    return 0 if status == "APROVADO" else 1


if __name__ == "__main__":
    raise SystemExit(main())
