"""Conta documentos da 00_landing e compara com A-OS_2022.xlsx.

Execute da raiz do projeto: python -u src/contar_arquivos_landing_os.py .
Aceita arquivos .xlsx e .xlsm; ignora temporários do Excel (~$...).
Não altera as planilhas. Códigos repetidos são alertas para conferir na OS.
"""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

EXTENSOES = {".xlsx", ".xlsm"}
CODIGO_OS = re.compile(r"OSMMQEQ\s*[-_. ]?\s*(\d+)\s*F\s*(\d+)", re.I)


def contar_linhas_anual(caminho: Path) -> tuple[int, int, int, dict[str, int]]:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        cabecalho = next(rows, ())
        coluna_os = next(
            (i for i, v in enumerate(cabecalho) if str(v).strip().casefold() == "formulário"),
            None,
        )
        if coluna_os is None:
            raise ValueError("Coluna 'formulário' ausente no arquivo anual")
        coluna_mes = next((i for i, v in enumerate(cabecalho) if str(v).strip().casefold() in {"mês", "mes"}), None)
        if coluna_mes is None:
            raise ValueError("Coluna 'mês' ausente no arquivo anual")
        linhas = 0
        vazios = 0
        codigos = set()
        meses = Counter()
        for row in rows:
            if not any(v is not None and str(v).strip() for v in row):
                continue
            linhas += 1
            meses[str(row[coluna_mes]).strip().upper()] += 1
            valor = row[coluna_os] if coluna_os < len(row) else None
            if valor is None or not str(valor).strip():
                vazios += 1
            else:
                codigos.add(str(valor).strip().upper())
        return linhas, len(codigos), vazios, dict(meses)
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("projeto", nargs="?", type=Path, default=Path("."))
    parser.add_argument("--run-id", default=None, help="Identificador da execução compartilhado com o orquestrador")
    args = parser.parse_args()
    raiz = args.projeto.resolve()
    landing = raiz / "datalake" / "00_landing"
    anual = raiz / "datalake" / "01_bronze_raw" / "A-OS_2022.xlsx"
    if not landing.is_dir():
        parser.error(f"Pasta não encontrada: {landing}")

    nomes = defaultdict(list)
    codigos = defaultdict(list)
    sem_codigo = []
    temporarios = []
    invalidos = []
    total = 0
    por_mes = []
    print("CONTAGEM DOS DOCUMENTOS DA 00_landing")
    print("Mês | .xlsx | .xlsm | Total")
    for pasta in sorted((p for p in landing.iterdir() if p.is_dir()), key=lambda p: p.name):
        arquivos = sorted(
            (p for p in pasta.rglob("*") if p.is_file() and p.suffix.lower() in EXTENSOES),
            key=lambda p: str(p).casefold(),
        )
        temporarios.extend(p for p in arquivos if p.name.startswith("~$"))
        candidatos = [p for p in arquivos if not p.name.startswith("~$")]
        validos = []
        for arq in candidatos:
            if zipfile.is_zipfile(arq):
                validos.append(arq)
            else:
                invalidos.append(arq.relative_to(landing))
        xlsx = sum(p.suffix.lower() == ".xlsx" for p in validos)
        xlsm = len(validos) - xlsx
        total += len(validos)
        por_mes.append({"pasta": pasta.name, "xlsx": xlsx, "xlsm": xlsm, "total": len(validos)})
        print(f"{pasta.name} | {xlsx} | {xlsm} | {len(validos)}")
        for arq in validos:
            relativo = arq.relative_to(landing)
            nomes[arq.name.casefold()].append(relativo)
            m = CODIGO_OS.search(arq.stem)
            if m:
                codigos[f"OSMMQEQ{int(m[1])}F{int(m[2])}"].append(relativo)
            else:
                sem_codigo.append(relativo)

    print(f"\nTotal de documentos válidos: {total}")
    print(f"Temporários (~$) ignorados: {len(temporarios)}")
    print(f"Arquivos .xlsx/.xlsm inválidos ignorados: {len(invalidos)}")
    for caminho in invalidos:
        print(f"  {caminho}")
    print("Abrindo o consolidado anual para comparação...", flush=True)
    erro_anual = None
    try:
        linhas, os_distintas, os_vazias, meses_anual = contar_linhas_anual(anual)
        print(f"Linhas de dados no anual: {linhas}")
        print(f"Diferença (documentos - linhas): {total - linhas:+d}")
        print(f"OS distintas no anual: {os_distintas}; linhas sem formulário: {os_vazias}")
    except (OSError, ValueError, zipfile.BadZipFile) as exc:
        linhas = os_distintas = os_vazias = None
        meses_anual = {}
        erro_anual = f"{type(exc).__name__}: {exc}"
        print(f"Não foi possível ler {anual}: {erro_anual}", flush=True)
    print("\nNOMES DE ARQUIVO REPETIDOS (entre pastas)")
    duplicados_nomes = {n: ps for n, ps in nomes.items() if len(ps) > 1}
    for nome, caminhos in sorted(duplicados_nomes.items()):
        print(f"{nome}: {len(caminhos)}")
        for caminho in caminhos:
            print(f"  {caminho}")
    print(f"Grupos: {len(duplicados_nomes)}; arquivos envolvidos: {sum(map(len, duplicados_nomes.values()))}")
    print("\nCÓDIGOS DE OS REPETIDOS NO NOME (inclusive variantes como '(2)')")
    duplicados_os = {codigo: ps for codigo, ps in codigos.items() if len(ps) > 1}
    for codigo, caminhos in sorted(duplicados_os.items()):
        print(f"{codigo}: {len(caminhos)}")
        for caminho in caminhos:
            print(f"  {caminho}")
    print(f"Códigos repetidos: {len(duplicados_os)}; arquivos envolvidos: {sum(map(len, duplicados_os.values()))}")
    print(f"Arquivos sem código OSMMQEQ... no nome: {len(sem_codigo)}")
    for caminho in sem_codigo:
        print(f"  {caminho}")
    print("\nConclusão: a igualdade de contagens é uma verificação inicial; códigos repetidos e")
    print("itens múltiplos devem ser conferidos na OS antes de declarar equivalência linha a linha.")
    aprovado = erro_anual is None and total == linhas and not sem_codigo
    run_id = args.run_id or datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    saida = raiz / "docs" / "execucoes" / "AB00" / run_id
    saida.mkdir(parents=True, exist_ok=False)
    relatorio = {
        "etapa": "AB00", "run_id": run_id,
        "status": "APROVADO" if aprovado else "REPROVADO",
        "contagem_por_mes": por_mes,
        "documentos_validos": total, "linhas_anual": linhas,
        "diferenca_documentos_menos_linhas": total - linhas if linhas is not None else None,
        "erro_leitura_anual": erro_anual,
        "os_distintas_anual": os_distintas, "os_vazias_anual": os_vazias,
        "linhas_anual_por_mes": meses_anual,
        "arquivos_invalidos": list(map(str, invalidos)),
        "arquivos_temporarios": [str(p.relative_to(landing)) for p in temporarios],
        "nomes_repetidos": {k: list(map(str, v)) for k, v in duplicados_nomes.items()},
        "codigos_repetidos": {k: list(map(str, v)) for k, v in duplicados_os.items()},
        "arquivos_sem_codigo": list(map(str, sem_codigo)),
        "criterio": "Comparação de contagens; não prova equivalência de registros ou itens.",
    }
    destino = saida / "conclusao_AB00.json"
    destino.write_text(json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")
    linhas_md = ["# AB00 — contagem da origem", "", f"Execução: `{run_id}`", "",
                 f"Status: **{relatorio['status']}**", "", "| Mês | XLSX | XLSM | Documentos | Linhas anual | Diferença |",
                 "|---|---:|---:|---:|---:|---:|"]
    for mes in por_mes:
        nome = mes["pasta"].split("_", 1)[1].rsplit("_", 1)[0]
        chave = nome.upper()
        anual_mes = meses_anual.get(chave)
        linhas_md.append(f"| {nome} | {mes['xlsx']} | {mes['xlsm']} | {mes['total']} | "
                         f"{anual_mes if anual_mes is not None else '—'} | "
                         f"{mes['total'] - anual_mes if anual_mes is not None else '—'} |")
    linhas_md += ["", f"Total válido na origem: **{total}**.",
                  f"Linhas no anual: **{linhas if linhas is not None else 'indisponível'}**.",
                  f"Arquivos inválidos ignorados: **{len(invalidos)}**.",
                  f"Códigos repetidos nos nomes: **{len(duplicados_os)}**.", "",
                  "A diferença de contagem requer conferência por documento; "
                  "uma OS pode ter múltiplos itens ou desdobramentos no anual.", "",
                  f"Detalhes e caminhos dos arquivos: `conclusao_AB00.json`."]
    (saida / "conclusao_AB00.md").write_text("\n".join(linhas_md) + "\n", encoding="utf-8")
    print(f"Evidência JSON: {destino}")
    return 0 if aprovado else 1


if __name__ == "__main__":
    raise SystemExit(main())
