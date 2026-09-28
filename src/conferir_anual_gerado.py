"""AB02: reconcilia A-OS_2022_gerado.xlsx com cada documento da 00_landing.

Da raiz do projeto: python -u .\\src\\conferir_anual_gerado.py . --origem-run-id 20260928_003359_127063
A saída é exclusiva de cada execução em docs/execucoes/AB02/<run_id>/.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

MESES = ("JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO")
COLUNAS = ("mês", "formulário", "prestador solicitado", "tipo", "número de controle",
           "descrição do defeito", "área/equipamento", "local", "descrição do serviço realizado",
           "peças e equipamentos necessários", "qtd", "valor unit", "valor total")
CODIGO = re.compile(r"OSM\.MQ/EQ-(\d+)F2", re.I)
NOME_CODIGO = re.compile(r"OSMMQEQ(\d+)F2", re.I)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def numero(valor: object) -> Decimal:
    if valor is None or str(valor).strip() == "":
        raise ValueError("valor numérico ausente")
    try:
        return Decimal(str(valor).replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"valor numérico inválido: {valor!r}") from exc


def normalizar(valor: object) -> str:
    return " ".join(str(valor or "").split()).casefold()


def codigo(valor: object) -> str | None:
    m = CODIGO.search(str(valor or ""))
    return f"OSM.MQ/EQ-{m.group(1)}F2" if m else None


def auditar_origem(path: Path, mes: str) -> tuple[dict, list[dict]]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if "OSM" not in wb:
            raise ValueError("aba OSM ausente")
        ws = wb["OSM"]
        os_codigo = codigo(ws["H2"].value)
        if not os_codigo:
            raise ValueError(f"código inválido em H2: {ws['H2'].value!r}")
        itens = []
        problemas = []
        for row in ws.iter_rows(min_row=36, max_row=60, min_col=1, max_col=10):
            lin = row[0].row
            qtd, peca, unit, total = row[0].value, row[1].value, row[7].value, row[9].value
            if all(v in (None, "") for v in (qtd, peca, unit, total)):
                continue
            if total in (None, "", 0):
                # A coluna A também contém rótulos de horário/observações após o quadro.
                if unit not in (None, "") or isinstance(qtd, (int, float, Decimal)):
                    problemas.append({"linha": lin, "erro": "item com quantidade/preço sem total"})
                continue
            try:
                q, u, t = numero(qtd), numero(unit), numero(total)
                if abs(q * u - t) > Decimal("0.005"):
                    problemas.append({"linha": lin, "erro": "quantidade x preço diverge do total",
                                      "qtd": str(q), "unit": str(u), "total": str(t)})
                itens.append({"linha": lin, "peca": peca, "qtd": q, "unit": u, "total": t})
            except ValueError as exc:
                problemas.append({"linha": lin, "erro": str(exc)})
        if not itens:
            problemas.append({"erro": "nenhum item precificado nas linhas 36 a 60"})
        nome = NOME_CODIGO.search(path.name)
        if nome and nome.group(1) != os_codigo.split("-")[1][:-2]:
            problemas.append({"erro": "código do nome diverge do formulário H2",
                              "nome": nome.group(1), "formulario": os_codigo})
        dados = dict(zip(COLUNAS[:9], [mes, os_codigo, ws["A19"].value, ws["A24"].value,
                      ws["A26"].value, ws["C19"].value, ws["F19"].value,
                      ws["H19"].value, ws["A31"].value]))
        dados[COLUNAS[9]] = itens[0]["peca"] if len(itens) == 1 else " | ".join(
            str(i["peca"] or f"ITEM LINHA {i['linha']}") for i in itens)
        dados[COLUNAS[10]] = itens[0]["qtd"] if len(itens) == 1 else None
        dados[COLUNAS[11]] = itens[0]["unit"] if len(itens) == 1 else None
        dados[COLUNAS[12]] = sum((i["total"] for i in itens), Decimal(0))
        return dados, problemas
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("projeto", nargs="?", type=Path, default=Path("."))
    parser.add_argument("--origem-run-id", default="20260928_003359_127063",
                        help="execução AB01 onde está a planilha gerada")
    parser.add_argument("--run-id", help="identificador da execução AB02; único")
    args = parser.parse_args()
    raiz = args.projeto.resolve()
    arquivo = raiz / "quality" / "AB01" / args.origem_run_id / "A-OS_2022_gerado.xlsx"
    landing = raiz / "datalake" / "00_landing"
    if not landing.is_dir() or not arquivo.is_file():
        parser.error(f"Verifique 00_landing e planilha gerada: {landing} | {arquivo}")
    run_id = args.run_id or datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    docs = raiz / "docs" / "execucoes" / "AB02" / run_id
    docs.mkdir(parents=True, exist_ok=False)
    erros = []
    vistos = defaultdict(list)
    fontes = {}
    contagem = Counter()
    arquivos_hash = {}
    for pasta in sorted(landing.iterdir()):
        if not pasta.is_dir() or not pasta.name[:2].isdigit() or not 1 <= int(pasta.name[:2]) <= 12:
            continue
        mes = MESES[int(pasta.name[:2]) - 1]
        for source in sorted(pasta.rglob("*")):
            if not source.is_file() or source.suffix.lower() not in {".xlsm", ".xlsx"} or source.name.startswith("~$"):
                continue
            rel = source.relative_to(raiz).as_posix()
            if not zipfile.is_zipfile(source):
                erros.append({"arquivo": rel, "erro": "arquivo Excel inválido"})
                continue
            try:
                arquivos_hash[rel] = sha256(source)
                dados, problemas = auditar_origem(source, mes)
                chave = dados["formulário"]
                vistos[chave].append(rel)
                fontes[chave] = (dados, rel)
                contagem[mes] += 1
                for problema in problemas:
                    erros.append({"arquivo": rel, **problema})
            except Exception as exc:
                erros.append({"arquivo": rel, "erro": f"falha de leitura: {type(exc).__name__}: {exc}"})
        print(f"[AB02] {mes}: {contagem[mes]} origens lidas", flush=True)
    for chave, paths in vistos.items():
        if len(paths) > 1:
            erros.append({"formulario": chave, "erro": "código repetido na origem", "arquivos": paths})
    workbook = load_workbook(arquivo, read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.sheetnames) != 1 or workbook.active.title != "OS_2022":
            erros.append({"erro": "aba(s) divergentes", "abas": workbook.sheetnames})
        ws = workbook.active
        cab = tuple(c.value for c in ws[1])
        if cab != COLUNAS:
            erros.append({"erro": "colunas divergentes", "esperado": COLUNAS, "obtido": cab})
        linhas = list(ws.iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    por_mes = Counter()
    linhas_codigo = defaultdict(list)
    for excel_row, linha in enumerate(linhas, 2):
        if not any(v is not None for v in linha):
            erros.append({"linha_excel": excel_row, "erro": "linha vazia"})
            continue
        chave = codigo(linha[1] if len(linha) > 1 else None)
        if not chave or normalizar(linha[1]) != normalizar(chave):
            erros.append({"linha_excel": excel_row, "erro": "formulário inválido", "valor": str(linha[1] if len(linha)>1 else None)})
            continue
        linhas_codigo[chave].append(excel_row)
        por_mes[str(linha[0]).upper()] += 1
        if chave not in fontes:
            erros.append({"linha_excel": excel_row, "formulario": chave, "erro": "O.S. sem documento correspondente"})
            continue
        esperado, rel = fontes[chave]
        for idx, campo in enumerate(COLUNAS):
            atual = linha[idx] if idx < len(linha) else None
            correto = esperado[campo]
            if idx in (10, 11, 12):
                try:
                    igual = (atual in (None, "") and correto is None) or abs(numero(atual) - numero(correto)) <= Decimal("0.005")
                except ValueError:
                    igual = False
            else:
                igual = normalizar(atual) == normalizar(correto)
            if not igual:
                erros.append({"linha_excel": excel_row, "formulario": chave, "arquivo_origem": rel,
                              "campo": campo, "esperado": str(correto), "obtido": str(atual)})
    for chave, (esperado, rel) in fontes.items():
        if chave not in linhas_codigo:
            erros.append({"formulario": chave, "arquivo_origem": rel, "erro": "O.S. ausente no arquivo gerado"})
    for chave, rows in linhas_codigo.items():
        if len(rows) > 1:
            erros.append({"formulario": chave, "linhas_excel": rows, "erro": "chave repetida no arquivo gerado"})
    for mes in MESES:
        if por_mes[mes] != contagem[mes]:
            erros.append({"mes": mes, "erro": "quantidade mensal divergente",
                          "landing": contagem[mes], "gerado": por_mes[mes]})
    if len(linhas) != sum(contagem.values()):
        erros.append({"erro": "quantidade total divergente", "landing": sum(contagem.values()), "gerado": len(linhas)})
    registro = {"etapa": "AB02", "run_id": run_id, "origem_run_id": args.origem_run_id,
                "data_hora": datetime.now().astimezone().isoformat(),
                "status": "APROVADO" if not erros else "REPROVADO",
                "criterio": "uma linha por documento; chave O.S. única; 13 campos e valores reconciliados com OSM da 00_landing",
                "planilha": arquivo.relative_to(raiz).as_posix(), "planilha_sha256": sha256(arquivo),
                "fontes_sha256": arquivos_hash, "documentos_landing": sum(contagem.values()),
                "linhas_geradas": len(linhas), "contagem_landing_por_mes": dict(contagem),
                "contagem_gerado_por_mes": dict(por_mes), "os_distintas_gerado": len(linhas_codigo),
                "erros": erros}
    (docs / "conclusao_AB02.json").write_text(json.dumps(registro, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    resumo = ["# AB02 — conferência do anual gerado", "", f"Execução: `{run_id}`", f"AB01 origem: `{args.origem_run_id}`",
              f"Status: **{registro['status']}**", "", f"Arquivos de origem: **{sum(contagem.values())}**; linhas do gerado: **{len(linhas)}**; O.S. distintas: **{len(linhas_codigo)}**.",
              f"Divergências: **{len(erros)}**.", "", "| Mês | Origem | Gerado |", "| --- | ---: | ---: |"]
    resumo += [f"| {m} | {contagem[m]} | {por_mes[m]} |" for m in MESES]
    if erros:
        resumo += ["", "## Divergências", ""] + [f"- `{json.dumps(e, ensure_ascii=False, default=str)}`" for e in erros[:100]]
        if len(erros) > 100:
            resumo.append(f"- Outras {len(erros)-100} no JSON.")
    resumo += ["", "A conferência cobre a equivalência dos campos mapeados da aba OSM e o quadro de itens A36:J60; não confere campos não usados desse documento."]
    (docs / "conclusao_AB02.md").write_text("\n".join(resumo) + "\n", encoding="utf-8")
    print(f"[AB02] {registro['status']}: {len(linhas)} linhas, {len(erros)} divergências. Evidência: {docs}", flush=True)
    return 0 if not erros else 1


if __name__ == "__main__":
    raise SystemExit(main())
