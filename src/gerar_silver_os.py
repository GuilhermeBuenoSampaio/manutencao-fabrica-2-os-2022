"""AB04: materializa uma O.S. por linha na Silver Parquet após AB03 aprovada.

Da raiz: python -u .\\src\\gerar_silver_os.py . --origem-run-id ID_AB01 --ab03-run-id ID_AB03
Requer openpyxl e pyarrow (python -m pip install openpyxl pyarrow).
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

MESES = ("JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO",
         "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO")
COLUNAS = ("mês", "formulário", "prestador solicitado", "tipo", "número de controle",
           "descrição do defeito", "área/equipamento", "setor", "local",
           "descrição do serviço realizado", "peças e equipamentos necessários",
           "qtd", "valor unit", "valor total")
OS = re.compile(r"^OSM\.MQ/EQ-\d+F2$", re.I)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as s:
        for chunk in iter(lambda: s.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def decimal(valor: object, casas: str, campo: str, codigo: str) -> Decimal:
    try:
        n = Decimal(str(valor))
        if not n.is_finite():
            raise ValueError()
        quant = n.quantize(Decimal(casas))
        if quant != n:
            raise ValueError()
        return quant
    except (InvalidOperation, ValueError):
        raise ValueError(f"{codigo}: {campo} inválido: {valor!r}") from None


def texto(valor: object) -> str | None:
    s = str(valor).strip() if valor is not None else ""
    return s or None


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("projeto", nargs="?", type=Path, default=Path("."))
    p.add_argument("--origem-run-id", default="20260928_003359_127063")
    p.add_argument("--ab03-run-id", help="execução AB03 aprovada; padrão: aprovação mais recente para os mesmos bytes")
    p.add_argument("--manual-id", default="manual-20260928_003359_127063")
    p.add_argument("--run-id")
    args = p.parse_args()
    raiz = args.projeto.resolve()
    original = raiz / "quality" / "AB01" / args.origem_run_id / "A-OS_2022_gerado.xlsx"
    manual = raiz / "quality" / "AB02" / args.manual_id / "classificacao_setor_os_2022.xlsx"
    if not original.is_file() or not manual.is_file():
        p.error(f"Entradas ausentes: {original} | {manual}")
    h_original, h_manual = sha256(original), sha256(manual)
    docs_ab03 = raiz / "docs" / "execucoes" / "AB03"
    candidatos = ([docs_ab03 / args.ab03_run_id / "conclusao_AB03.json"] if args.ab03_run_id else
                  sorted(docs_ab03.glob("*/conclusao_AB03.json"), reverse=True) if docs_ab03.exists() else [])
    aprovado = None
    for path in candidatos:
        try:
            r = json.loads(path.read_text(encoding="utf-8"))
            if (r.get("status") == "APROVADO" and r.get("sha256_origem") == h_original and
                r.get("sha256_manual") == h_manual and r.get("linhas_origem") == r.get("linhas_manual") and
                not r.get("divergencias")):
                aprovado = path
                break
        except (OSError, ValueError):
            continue
    if aprovado is None:
        p.error("AB03 aprovada para os bytes atuais dos dois arquivos não encontrada. Execute AB03 novamente.")
    run_id = args.run_id or datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    destino = raiz / "datalake" / "02_silver" / run_id
    docs = raiz / "docs" / "execucoes" / "AB04" / run_id
    if destino.exists() or docs.exists():
        p.error(f"ID já existe: {run_id}")
    try:
        import pyarrow as pa
        import pyarrow.parquet as pq
    except ImportError:
        p.error("Instale pyarrow no interpretador usado: python -m pip install pyarrow")
    livro = load_workbook(manual, read_only=True, data_only=True, keep_links=False)
    linhas = []
    try:
        if livro.sheetnames != ["OS_2022"]:
            p.error(f"Abas inesperadas: {livro.sheetnames}")
        rows = livro.active.iter_rows(values_only=True)
        header = tuple(next(rows))
        if header != COLUNAS:
            p.error(f"Cabeçalho incorreto: {header}")
        vistos = set()
        for n, row in enumerate(rows, 2):
            if len(row) != len(COLUNAS):
                raise ValueError(f"Linha {n}: quantidade de colunas divergente")
            x = dict(zip(COLUNAS, row))
            codigo = texto(x["formulário"])
            if not codigo or not OS.fullmatch(codigo) or codigo in vistos:
                raise ValueError(f"Linha {n}: O.S. inválida ou duplicada: {codigo!r}")
            vistos.add(codigo)
            mes = texto(x["mês"])
            if mes not in MESES:
                raise ValueError(f"{codigo}: mês inválido: {mes!r}")
            setor = texto(x["setor"])
            if setor not in ("empanados", "pão de queijo", "geral"):
                raise ValueError(f"{codigo}: setor inválido: {setor!r}")
            qtd = (None if x["qtd"] in (None, "") else decimal(x["qtd"], "0.001", "qtd", codigo))
            unit = (None if x["valor unit"] in (None, "") else decimal(x["valor unit"], "0.01", "valor unit", codigo))
            total = decimal(x["valor total"], "0.01", "valor total", codigo)
            if (qtd is None) != (unit is None):
                raise ValueError(f"{codigo}: quantidade/preço parcialmente ausentes")
            if qtd is not None and abs(qtd * unit - total) > Decimal("0.005"):
                raise ValueError(f"{codigo}: quantidade x preço diverge do total")
            linhas.append({"formulario_os": codigo, "ano": 2022, "mes_numero": MESES.index(mes) + 1,
                          "mes": mes, "setor": setor,
                          "prestador_solicitado": texto(x["prestador solicitado"]),
                          "tipo": texto(x["tipo"]),
                          "numero_controle": texto(x["número de controle"]),
                          "descricao_defeito": texto(x["descrição do defeito"]),
                          "area_equipamento": texto(x["área/equipamento"]),
                          "local": texto(x["local"]),
                          "descricao_servico_realizado": texto(x["descrição do serviço realizado"]),
                          "pecas_equipamentos_necessarios": texto(x["peças e equipamentos necessários"]),
                          "qtd": qtd, "valor_unitario_brl": unit, "valor_total_brl": total})
    finally:
        livro.close()
    ab03 = json.loads(aprovado.read_text(encoding="utf-8"))
    if len(linhas) != ab03["linhas_manual"]:
        raise ValueError("Contagem difere da AB03 aprovada")
    schema = pa.schema([
        ("formulario_os", pa.string()), ("ano", pa.int16()), ("mes_numero", pa.int8()),
        ("mes", pa.string()), ("setor", pa.string()), ("prestador_solicitado", pa.string()),
        ("tipo", pa.string()), ("numero_controle", pa.string()), ("descricao_defeito", pa.string()),
        ("area_equipamento", pa.string()), ("local", pa.string()),
        ("descricao_servico_realizado", pa.string()), ("pecas_equipamentos_necessarios", pa.string()),
        ("qtd", pa.decimal128(18, 3)), ("valor_unitario_brl", pa.decimal128(18, 2)),
        ("valor_total_brl", pa.decimal128(18, 2))])
    tabela = pa.Table.from_pylist(linhas, schema=schema)
    destino.mkdir(parents=True, exist_ok=False)
    path = destino / "os_2022_silver.parquet"
    pq.write_table(tabela, path, compression="snappy")
    lido = pq.read_table(path)
    if lido.num_rows != len(linhas) or lido.schema != schema or lido.to_pylist() != tabela.to_pylist():
        raise RuntimeError("Parquet gravado diverge dos registros de entrada")
    docs.mkdir(parents=True, exist_ok=False)
    setores = Counter(v["setor"] for v in linhas)
    meses = Counter(v["mes"] for v in linhas)
    total = sum((v["valor_total_brl"] for v in linhas), Decimal(0))
    report = {"etapa": "AB04", "run_id": run_id, "status": "APROVADO",
              "granularidade": "uma linha por formulário/O.S.", "chave_primaria": "formulario_os",
              "linhas": len(linhas), "os_distintas": len({v['formulario_os'] for v in linhas}),
              "cardinalidade_setor": dict(sorted(setores.items())), "linhas_por_mes": dict(meses),
              "custo_total_brl": str(total), "regras": ["mes_numero derivado do mês", "ano fixo 2022",
                  "texto aparado sem alterar a fonte", "custos em decimal de duas casas",
                  "qtd em decimal de três casas", "número de controle mantido como texto",
                  "qtd e valor unitário podem ser nulos juntos para O.S. com vários itens"],
              "origem": original.relative_to(raiz).as_posix(), "sha256_origem": h_original,
              "classificacao": manual.relative_to(raiz).as_posix(), "sha256_classificacao": h_manual,
              "aprovacao_ab03": aprovado.relative_to(raiz).as_posix(), "parquet": path.relative_to(raiz).as_posix(),
              "sha256_parquet": sha256(path), "schema": str(schema),
              "limite": "A classificação de setor foi feita manualmente; despesas de parada de produção não constam da fonte."}
    (docs / "conclusao_AB04.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    md = ["# AB04 — Silver de ordens de serviço", "", f"Execução: `{run_id}`; status: **APROVADO**.",
          f"Linhas: **{len(linhas)}**; O.S. distintas: **{report['os_distintas']}**; custo registrado: **R$ {total}**.",
          f"Parquet: `{report['parquet']}` (SHA-256 `{report['sha256_parquet']}`).",
          f"Classificação manual: `{report['classificacao']}`; validação AB03: `{report['aprovacao_ab03']}`.",
          "", "## Cardinalidade por setor", ""]
    md += [f"- {k}: {v} O.S." for k, v in sorted(setores.items())]
    md += ["", "Chave primária: `formulario_os`. Granularidade: uma O.S. por linha.",
           "Campos de quantidade e preço unitário podem estar vazios juntos quando uma O.S. tem vários itens; o total conserva a soma.",
           report["limite"]]
    (docs / "conclusao_AB04.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"[AB04] APROVADO: {len(linhas)} O.S.; R$ {total}; Parquet: {path}; docs: {docs}", flush=True)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValueError, RuntimeError) as exc:
        print(f"[AB04] REPROVADO: {exc}", file=sys.stderr)
        raise SystemExit(1) from None
